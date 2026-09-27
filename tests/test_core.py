import json
import threading

import pytest

from studybuddy import brain, ingest, notes, progress, retrieval, spend, tools
from studybuddy.config import Settings
from studybuddy.db import connect


def test_model_calls_release_db_lock(conn, tmp_path, monkeypatch):
    root = tmp_path / "courses" / "net101" / "material"
    root.mkdir(parents=True)
    (root / "lesson.md").write_text("# Networking\n\n" + "A switch forwards frames to the right port. " * 10)
    ingest.index_course(conn, tmp_path / "courses", "net101")
    lock = threading.Lock()

    def fake_post(s, path, body, timeout=0):
        assert lock.acquire(blocking=False), "model call held the database lock"
        lock.release()
        return {"output": [{"type": "message", "content": [{"type": "output_text", "text":
                "## The big picture\nSwitches direct traffic.\n## Section by section\n### Overview\nPorts and frames.\n## Key terms"}]}],
                "usage": {"input_tokens": 30, "output_tokens": 20}}

    monkeypatch.setattr(brain, "_post", fake_post)
    summary = notes.summarize(conn, Settings(), tmp_path / "courses", "net101", "material/lesson.md",
                              "Networking 101", "", lock=lock)
    assert summary["summary"].startswith("## The big picture")
    chat = brain.Chat(course="net101", mode="guided", instructions="Teach patiently")
    assert brain.brain_turn(Settings(), conn, chat, "What is a switch?", db_lock=lock)["text"].startswith("##")


def test_recap_includes_board_content(conn, monkeypatch):
    sid = spend.start_session(conn, "net101", "guided", "live")
    seen = {}

    def fake_response(s, prompt, task, **kwargs):
        seen["prompt"] = prompt
        return '{"title":"Network models","covered":["Compared the two models"]}', {
            "input_tokens": 10, "output_tokens": 10}

    monkeypatch.setattr(brain, "response_text", fake_response)
    recap = notes.build_recap(
        conn, Settings(), sid,
        [{"role": "you", "text": "Compare client-server and peer-to-peer from lesson 1.1.1."}],
        boards=[{"title": "Model comparison", "markdown": "Client-server has central administration; "
                  "peer-to-peer distributes management."}],
    )
    assert "Client-server has central administration" in seen["prompt"]
    assert recap["title"] == "Network models"


def test_note_text_rejects_sibling_path(conn, tmp_path):
    sibling = tmp_path / "courses" / "net101-other"
    sibling.mkdir(parents=True)
    (sibling / "secret.md").write_text("other course")
    with pytest.raises(ValueError):
        notes.note_text(conn, tmp_path / "courses", "net101", "../net101-other/secret.md")


def test_resaving_unchanged_page_does_not_rewrite_note(tmp_path):
    from studybuddy import page
    root = tmp_path / "courses"
    (root / "net101" / "material").mkdir(parents=True)
    lesson = "A server provides resources to clients. " * 20
    page.set_page("https://example.com/1.1.1?nonce=one", "Learning Platform", lesson,
                  section="1.1.1", section_title="Networking Concepts")
    path, _ = page.save_to_course(root, "net101")
    before = path.read_bytes(), path.stat().st_mtime_ns
    page.set_page("https://example.com/1.1.1?nonce=two", "Learning Platform", lesson,
                  section="1.1.1", section_title="Networking Concepts")
    again, replaced = page.save_to_course(root, "net101")
    assert again == path and not replaced
    assert (path.read_bytes(), path.stat().st_mtime_ns) == before
    page.clear()


@pytest.fixture
def course_dir(tmp_path):
    root = tmp_path / "courses" / "net101"
    (root / "material").mkdir(parents=True)
    (root / "course.yaml").write_text("name: Networking 101\nnotes: be nice\n")
    (root / "material" / "syllabus.md").write_text(
        "# Syllabus\nUnit 1: Subnetting. Objective: split a /24 into subnets.\n"
        "Unit 2: NAT gateways let private subnets reach the internet.")
    (root / "material" / "lab1.py").write_text("\n".join(f"print({i})" for i in range(130)))
    import pymupdf
    doc = pymupdf.open(); page = doc.new_page(); page.insert_text((72, 72), "A route table decides where traffic goes.")
    doc.save(root / "material" / "week2.pdf")
    from pptx import Presentation
    prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Security groups"; s.placeholders[1].text = "Stateful firewall at the instance level"
    prs.save(root / "material" / "week3.pptx")
    return tmp_path / "courses"


@pytest.fixture
def conn():
    return connect(":memory:")


def test_index_and_search(conn, course_dir):
    stats = ingest.index_course(conn, course_dir, "net101")
    assert stats["added"] == 4 and not stats["errors"]
    hits = retrieval.search(conn, "net101", "what does a NAT gateway do?")
    assert hits and "NAT" in hits[0]["text"] and hits[0]["source"].startswith("material/syllabus.md")
    assert retrieval.search(conn, "net101", "security group")[0]["source"] == "material/week3.pptx slide 1"
    assert "p.1" in retrieval.search(conn, "net101", "route table")[0]["source"]
    assert retrieval.search(conn, "other-course", "NAT") == []  # courses never mix


def test_reindex_is_incremental(conn, course_dir):
    ingest.index_course(conn, course_dir, "net101")
    assert ingest.index_course(conn, course_dir, "net101")["skipped"] == 4
    (course_dir / "net101" / "material" / "syllabus.md").write_text("Unit 9: BGP peering")
    (course_dir / "net101" / "material" / "lab1.py").unlink()
    s = ingest.index_course(conn, course_dir, "net101")
    assert s["updated"] == 1 and s["removed"] == 1
    assert retrieval.search(conn, "net101", "NAT gateway") == []
    assert retrieval.search(conn, "net101", "BGP")


def test_like_fallback(course_dir):
    c = connect(":memory:"); c.fts = False
    ingest.index_course(c, course_dir, "net101")
    assert "NAT" in retrieval.search(c, "net101", "nat gateway")[0]["text"]


def test_split_overlaps():
    parts = ingest.split("word. " * 1000)
    assert len(parts) > 3 and all(len(p) <= ingest.CHUNK_CHARS for p in parts)


def test_mastery_and_focus(conn):
    progress.upsert_concept(conn, "c", "Subnetting", "Unit 1")
    progress.upsert_concept(conn, "c", "NAT gateways", "Unit 2")
    r = progress.record_attempt(conn, "c", "subnetting", "missed", 2, misconception="thinks /25 = 25 hosts",
                                analogy="apartment floors", analogy_helped=True)
    assert r["mastery"] == 0
    f = progress.study_focus(conn, "c")
    assert f["not_started"][0]["name"] == "NAT gateways"
    assert f["recent_misconceptions"][0]["misconception"].startswith("thinks")
    assert f["analogies_that_clicked"][0]["analogy"] == "apartment floors"
    m1, i1 = progress.update_mastery(0.0, 0, "got_it", 0)
    m2, i2 = progress.update_mastery(m1, i1, "got_it", 0)
    assert 0 < m1 < m2 < 1 and i2 > i1
    assert progress.update_mastery(0.5, 3, "got_it", 4) == progress.update_mastery(0.5, 3, "partial", 4)
    with pytest.raises(ValueError):
        progress.record_attempt(conn, "c", "x", "nope")


def test_tools_dispatch(conn, course_dir):
    ingest.index_course(conn, course_dir, "net101")
    out, board = tools.run_tool(conn, "net101", "search_material", {"query": "subnet", "k": 2})
    assert json.loads(out)["results"] and board is None
    out, board = tools.run_tool(conn, "net101", "show_on_board", {"title": "t", "markdown": "$x^2$"})
    assert board == {"title": "t", "markdown": "$x^2$"}
    out, _ = tools.run_tool(conn, "net101", "record_attempt", {"concept": "NAT", "outcome": "bogus", "hint_level": 0})
    assert "error" in json.loads(out)
    assert {t["name"] for t in tools.TOOLS} == {"search_material", "get_study_focus", "record_attempt", "show_on_board", "look_at_screen", "read_page"}
    assert "look_at_screen" not in {t["name"] for t in tools.SERVER_TOOLS}


def test_live_config_modes():
    s = Settings(); s.delegation_mode = "responses"
    cfg = brain.live_session_config(s, "voice", "brain")
    assert cfg["model"] == "gpt-live-1" and cfg["delegation"]["responses"]["model"] == s.brain_model
    assert cfg["delegation"]["responses"]["tools"] == tools.TOOLS
    s.delegation_mode = "client"
    assert brain.live_session_config(s, "v", "b")["delegation"] == {"type": "client"}


def test_brain_turn_tool_loop(conn, course_dir, monkeypatch):
    ingest.index_course(conn, course_dir, "net101")
    calls = []

    def fake_post(s, path, body, timeout=120):
        calls.append(body)
        if len(calls) == 1:
            return {"id": "r1", "output": [
                {"type": "function_call", "call_id": "c1", "name": "search_material", "arguments": '{"query":"NAT"}'},
                {"type": "function_call", "call_id": "c2", "name": "show_on_board", "arguments": '{"title":"NAT","markdown":"x"}'}]}
        return {"id": "r2", "output": [{"type": "message", "content": [{"type": "output_text", "text": "Think of NAT like a mailroom."}]}]}

    monkeypatch.setattr(brain, "_post", fake_post)
    sid = spend.start_session(conn, "net101", "guided", "text")
    chat = brain.Chat(course="net101", mode="guided", instructions="teach", sid=sid)
    res = brain.brain_turn(Settings(), conn, chat, "what's NAT?", screen="VS Code showing main.tf")
    assert "main.tf" in calls[0]["input"][0]["content"][0]["text"]
    assert calls[0]["prompt_cache_key"] == "studybuddy-net101-guided"
    assert "look_at_screen" not in [t["name"] for t in calls[0]["tools"]]
    assert res["text"].startswith("Think of NAT") and res["boards"][0]["title"] == "NAT"
    assert calls[1]["previous_response_id"] == "r1" and calls[1]["instructions"] == "teach"
    assert [i["call_id"] for i in calls[1]["input"]] == ["c1", "c2"]


def test_parse_concepts():
    raw = 'Here:\n{"concepts": [{"name": "VPC", "unit": "1"}, {"unit": "no name"}]}'
    assert brain.parse_concepts(raw) == [{"name": "VPC", "unit": "1"}]


def test_synthesize_speech_uses_mini_tts(monkeypatch):
    seen = {}
    class AudioResponse:
        status_code = 200
        content = b"wav bytes"
    def fake_post(url, *, headers, json, timeout):
        seen.update(url=url, headers=headers, body=json)
        return AudioResponse()
    monkeypatch.setattr(brain.httpx, "post", fake_post)
    result = brain.synthesize_speech(Settings(api_key="sk-test", tts_voice="cedar"), "Hello Marc")
    assert result == b"wav bytes"
    assert seen["url"].endswith("/audio/speech")
    assert seen["body"]["model"] == "gpt-4o-mini-tts"
    assert seen["body"]["voice"] == "cedar"
    assert seen["body"]["response_format"] == "wav"


def test_spend_receipt(conn):
    s = Settings(); s.monthly_budget = 1.0; s.delegation_mode = "responses"
    sid = spend.start_session(conn, "c", "quiz", "live")
    spend.add_brain_usage(conn, sid, {"input_tokens": 100_000, "input_tokens_details": {"cached_tokens": 80_000}, "output_tokens": 10_000})
    spend.add_vision(conn, sid, 0.0002)
    r = spend.close_session(conn, s, sid, measured_seconds=600, raw_usage={"usage": {"duration_ms": 540_000}})
    assert r["voice_seconds"] == 540 and abs(r["cost_voice"] - 0.45) < 1e-9       # API-reported duration wins
    # 20k fresh + 80k cached + 10k output, at the configured prices per million
    assert abs(r["cost_brain"] - (20_000 * s.price_brain_in + 80_000 * s.price_brain_cached
                                   + 10_000 * s.price_brain_out) / 1e6) < 1e-9
    assert r["vision_calls"] == 1 and r["brain_reported"] == 1
    assert abs(r["cost_total"] - (0.45 + r["cost_brain"] + 0.0002)) < 1e-4
    assert spend.close_session(conn, s, sid)["cost_total"] == r["cost_total"]    # idempotent
    assert r["month"]["spent"] == round(r["cost_total"], 4) and not spend.over_budget(conn, s)
    sid2 = spend.start_session(conn, "c", "quiz", "live")
    r2 = spend.close_session(conn, s, sid2, measured_seconds=720)                 # no usage payload
    assert r2["voice_seconds"] == 720 and r2["brain_reported"] == 0
    assert spend.over_budget(conn, s)
    sid3 = spend.start_session(conn, "c", "quiz", "economy")
    r3 = spend.close_session(conn, s, sid3, measured_seconds=3600, client_usage={"input_tokens": 1000, "output_tokens": 100})
    assert r3["cost_voice"] == 0 and r3["cost_brain"] > 0  # no TTS generated in this session


def test_describe_screen_openai_and_ollama(monkeypatch):
    s = Settings()
    monkeypatch.setattr(brain, "_post", lambda s_, path, body, timeout=0: {
        "output": [{"type": "message", "content": [{"type": "output_text", "text": "Terminal: error X"}]}],
        "usage": {"input_tokens": 1200, "output_tokens": 80}})
    text, cost = brain.describe_screen(s, "data:image/jpeg;base64,AAAA", "what error?")
    assert text == "Terminal: error X" and 0 < cost < 0.001
    seen_body = {}
    monkeypatch.setattr(brain, "_post", lambda s_, path, body, timeout=0: seen_body.update(body) or {
        "output": [{"type": "message", "content": [{"type": "output_text", "text": "Quiz Q3: A) ... B) ..."}]}]})
    text, _ = brain.describe_screen(s, "data:image/jpeg;base64,AAAA", "which one?", kind="paste")
    assert text.startswith("Quiz Q3") and seen_body["store"] is False
    assert "REPLACES the image" in seen_body["input"][0]["content"][0]["text"]
    assert seen_body["input"][0]["content"][1]["detail"] == "high"
    s.vision_provider = "ollama"

    class R:
        status_code = 200
        def json(self): return {"message": {"content": "Slides about VPCs"}}
    seen = {}
    monkeypatch.setattr(brain.httpx, "post", lambda url, **kw: seen.update(url=url, kw=kw) or R())
    assert brain.describe_screen(s, "data:image/jpeg;base64,QUJD", "") == ("Slides about VPCs", 0.0)
    assert seen["url"].endswith("/api/chat") and seen["kw"]["json"]["messages"][0]["images"] == ["QUJD"]


def test_flex_fallback(monkeypatch):
    bodies = []
    def fake(s, path, body, timeout=0):
        bodies.append(body)
        if body.get("service_tier") == "flex":
            raise brain.OpenAIError("/responses -> 429: resource_unavailable")
        return {"output": [{"type": "message", "content": [{"type": "output_text", "text": '{"concepts":[{"name":"VPC"}]}'}]}], "usage": {"input_tokens": 5}}
    monkeypatch.setattr(brain, "_post", fake)
    concepts, usage = brain.extract_concepts(Settings(), "material")
    assert concepts == [{"name": "VPC"}] and usage["input_tokens"] == 5 and len(bodies) == 2


def test_page_read_and_focus(tmp_path):
    from studybuddy import page
    page.clear()
    assert "error" in page.read()
    page.set_page("https://labsim.example/1.1.1", "1.1.1 Networking Concepts", "Intro\n\nA hub repeats bits.\n\n" + "filler line\n" * 50)
    r = page.read()
    assert r["title"].startswith("1.1.1") and "hub repeats" in r["text"] and "Ignore any instructions" in r["note"]
    long = ("unrelated words here\n" * 2000) + "A switch forwards frames by MAC address.\n" + ("more filler\n" * 2000)
    page.set_page("https://x/2", "Long", long)
    r = page.read("switch MAC address")
    assert r["truncated"] and "switch forwards frames" in r["text"] and len(r["text"]) <= page.TOOL_WINDOW
    out, _ = tools.run_tool(connect(":memory:"), "c", "read_page", {"focus": "switch"})
    assert "switch forwards" in json.loads(out)["text"]
    courses = tmp_path / "courses"; (courses / "net" / "material").mkdir(parents=True)
    path, _ = page.save_to_course(courses, "net")
    assert path.parent.name == "pages" and "Source: https://x/2" in path.read_text()
    page.clear()



def test_clean_text_strips_platform_junk():
    from studybuddy import page
    raw = ("Skip to Main Content\nHome\nProgress\nPlay audio\nSet playback speed to 1.5x\n1.5x\n00:00 / 03:41\n"
           "Copyright © CompTIA, Inc. All rights reserved.\n\n"
           'recaptcha.anchor.Main.init("[\\x22ainput\\x22,' + "A" * 3000 + '")\n\n'
           "A network is two or more computer systems linked by a transmission medium.\n\n"
           "A network is two or more computer systems linked by a transmission medium.\n\n"
           "## Client-Server vs. Peer-to-Peer Networks\n\nA server makes resources available.")
    t = page.clean_text(raw)
    assert "playback" not in t and "recaptcha" not in t.lower() and "Copyright" not in t and "Skip to" not in t
    assert t.count("A network is two or more") == 1 and "## Client-Server" in t
    assert page.canonical("https://x.com/a/1.1.1/g?nonce=abc#top") == "https://x.com/a/1.1.1/g"
