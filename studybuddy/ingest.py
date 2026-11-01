"""Turn course files into searchable chunks.

Layout:  courses/<course-slug>/course.yaml      (name, notes for the tutor)
         courses/<course-slug>/material/...     (PDF, PPTX, DOCX, MD, TXT, code)
Re-running is cheap: unchanged files (same sha1) are skipped.
"""
from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Iterator

import yaml

from .db import Conn

CODE_EXT = {".py", ".js", ".ts", ".java", ".c", ".cpp", ".h", ".cs", ".go", ".rs", ".sql",
            ".sh", ".ps1", ".json", ".yaml", ".yml", ".tf", ".html", ".css"}
TEXT_EXT = {".md", ".txt", ".csv"}
SUPPORTED = CODE_EXT | TEXT_EXT | {".pdf", ".pptx", ".docx"}

CHUNK_CHARS = 1400
OVERLAP = 200


@dataclass
class Piece:
    locator: str
    text: str


# ---------- extraction ----------

def _pdf(path: Path) -> Iterator[Piece]:
    import pymupdf
    with pymupdf.open(path) as doc:
        for i, page in enumerate(doc, start=1):
            yield Piece(f"p.{i}", page.get_text())


def _pptx(path: Path) -> Iterator[Piece]:
    from pptx import Presentation
    prs = Presentation(str(path))
    for i, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
            if notes.strip():
                parts.append("Speaker notes: " + notes)
        yield Piece(f"slide {i}", "\n".join(parts))


def _docx(path: Path) -> Iterator[Piece]:
    from docx import Document
    doc = Document(str(path))
    heading, buf = "start", []
    for p in doc.paragraphs:
        if p.style.name.lower().startswith("heading") and p.text.strip():
            if buf:
                yield Piece(f"§ {heading}", "\n".join(buf))
            heading, buf = p.text.strip()[:60], [p.text]
        elif p.text.strip():
            buf.append(p.text)
    if buf:
        yield Piece(f"§ {heading}", "\n".join(buf))


def _code(path: Path) -> Iterator[Piece]:
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    step = 60
    for start in range(0, max(len(lines), 1), step - 10):
        block = lines[start:start + step]
        if not block:
            break
        yield Piece(f"lines {start + 1}-{start + len(block)}", "\n".join(block))
        if start + step >= len(lines):
            break


def _text(path: Path) -> Iterator[Piece]:
    yield Piece("", path.read_text(encoding="utf-8", errors="replace"))


def extract(path: Path) -> Iterator[Piece]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _pdf(path)
    if ext == ".pptx":
        return _pptx(path)
    if ext == ".docx":
        return _docx(path)
    if ext in CODE_EXT:
        return _code(path)
    return _text(path)


def split(text: str, size: int = CHUNK_CHARS, overlap: int = OVERLAP) -> list[str]:
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) <= size:
        return [text] if text else []
    out, start = [], 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):  # prefer to break on a paragraph / sentence
            cut = max(text.rfind("\n\n", start, end), text.rfind(". ", start, end))
            if cut > start + size // 2:
                end = cut + 1
        out.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return [c for c in out if c]


# ---------- courses ----------

@lru_cache(maxsize=128)
def _course_yaml(path: Path, modified_ns: int) -> dict:
    """Avoid parsing unchanged course metadata on every sidebar request."""
    return yaml.safe_load(path.read_text(encoding="utf-8")) or {}

def list_courses(courses_dir: Path) -> list[dict]:
    out = []
    if not courses_dir.exists():
        return out
    for d in sorted(p for p in courses_dir.iterdir() if p.is_dir() and not p.name.startswith(".")):
        meta = {}
        cfg = d / "course.yaml"
        if cfg.exists():
            meta = _course_yaml(cfg, cfg.stat().st_mtime_ns)
        if meta.get("hidden"):
            continue
        out.append({"slug": d.name, "name": str(meta.get("name", d.name)), "notes": meta.get("notes", ""),
                    "kind": meta.get("kind", ""), "code": str(meta.get("code", "") or ""),
                    "instructor": meta.get("instructor", ""), "term": meta.get("term", ""),
                    "order": meta.get("order", 99)})
    out.sort(key=lambda c: (c["order"], c["name"]))
    return out


def course_meta(courses_dir: Path, slug: str) -> dict:
    for c in list_courses(courses_dir):
        if c["slug"] == slug:
            return c
    raise KeyError(slug)


def _sha(path: Path) -> str:
    h = hashlib.sha1()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def _delete_file(conn: Conn, course: str, rel: str) -> None:
    ids = [r[0] for r in conn.execute("SELECT id FROM chunks WHERE course=? AND path=?", (course, rel))]
    if conn.fts:
        for i in ids:
            row = conn.execute("SELECT text FROM chunks WHERE id=?", (i,)).fetchone()
            conn.execute("INSERT INTO chunks_fts(chunks_fts, rowid, text) VALUES('delete', ?, ?)", (i, row[0]))
    conn.execute("DELETE FROM chunks WHERE course=? AND path=?", (course, rel))
    conn.execute("DELETE FROM files WHERE course=? AND path=?", (course, rel))


def index_course(conn: Conn, courses_dir: Path, slug: str) -> dict:
    """Index new/changed files in courses/<slug>/, drop deleted ones. Returns stats."""
    root = courses_dir / slug
    stats = {"added": 0, "updated": 0, "removed": 0, "skipped": 0, "errors": []}
    seen = set()
    for path in sorted(root.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED or path.name == "course.yaml":
            continue
        rel = path.relative_to(root).as_posix()
        seen.add(rel)
        try:
            sha = _sha(path)
        except OSError as e:
            stats["errors"].append(f"{rel}: {e}")
            continue
        row = conn.execute("SELECT sha FROM files WHERE course=? AND path=?", (slug, rel)).fetchone()
        if row and row["sha"] == sha:
            stats["skipped"] += 1
            continue
        try:
            pieces = list(extract(path))
        except Exception as e:  # corrupt/locked file: report, keep going
            stats["errors"].append(f"{rel}: {e}")
            continue
        if row:
            _delete_file(conn, slug, rel)
        for piece in pieces:
            for chunk in split(piece.text):
                cur = conn.execute("INSERT INTO chunks(course, path, locator, text) VALUES (?,?,?,?)",
                                   (slug, rel, piece.locator, chunk))
                if conn.fts:
                    conn.execute("INSERT INTO chunks_fts(rowid, text) VALUES (?,?)", (cur.lastrowid, chunk))
        conn.execute("INSERT OR REPLACE INTO files(course, path, sha, kind) VALUES (?,?,?,?)",
                     (slug, rel, sha, path.suffix.lower().lstrip(".")))
        stats["updated" if row else "added"] += 1
    for (rel,) in conn.execute("SELECT path FROM files WHERE course=?", (slug,)).fetchall():
        if rel not in seen:
            _delete_file(conn, slug, rel)
            stats["removed"] += 1
    conn.commit()
    return stats


def course_text(conn: Conn, slug: str, prefer: str = "syllabus", limit_chars: int = 400_000) -> str:
    """Concatenated course text for whole-course passes (e.g. building the concept map).
    Files whose path contains `prefer` come first."""
    rows = conn.execute(
        "SELECT path, locator, text FROM chunks WHERE course=? "
        "ORDER BY (instr(lower(path), ?) = 0), path, id", (slug, prefer)).fetchall()
    out, total = [], 0
    for r in rows:
        piece = f"[{r['path']} {r['locator']}]\n{r['text']}\n"
        total += len(piece)
        if total > limit_chars:
            break
        out.append(piece)
    return "\n".join(out)
